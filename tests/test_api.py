"""HTTP contracts with synthetic evidence; no Azure service calls."""

import csv
import io
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth import Principal, get_principal
from app.db import get_db
from app.main import app
from app.models import Base, Case, Document, Employee


@pytest.fixture
def api(monkeypatch):
    import app.auth as auth
    import app.main as main_module

    monkeypatch.setattr(main_module, "validate_configuration", lambda: None)
    # Configured production authentication must reject requests without tokens.
    monkeypatch.setattr(
        auth,
        "settings",
        SimpleNamespace(
            environment="production",
            allow_dev_auth=False,
            entra_tenant_id="tenant-a",
            entra_audience="test-api",
        ),
    )
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)

    def database():
        with Session(engine, expire_on_commit=False) as session:
            yield session

    actor = {"principal": Principal("alice", "tenant-a", frozenset())}
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[get_principal] = lambda: actor["principal"]
    with TestClient(app) as client:
        yield client, actor, engine
    app.dependency_overrides.clear()
    engine.dispose()


def create(client, department="Engineering"):
    response = client.post("/api/cases", json={"department": department, "consent": True})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def evidence(engine, case_id):
    with Session(engine) as db:
        ids = []
        for kind in ("pan", "aadhaar"):
            document = Document(
                case_id=case_id,
                tenant_id="tenant-a",
                doc_type=kind,
                filename="synthetic.pdf",
                content_type="application/pdf",
                storage_key=kind,
                sha256="a" * 64,
                scan_status="clean",
                extraction={"status": "complete"},
            )
            db.add(document)
            db.flush()
            ids.append(document.id)
        db.commit()
        return ids


def test_http_consent_authentication_and_tenant_isolation(api):
    client, actor, engine = api
    denied = client.post("/api/cases", json={"department": "Engineering", "consent": False})
    assert denied.status_code == 422
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(Case)) == 0
    case_id = create(client)
    actor["principal"] = Principal("bob", "tenant-a")
    assert client.get(f"/api/cases/{case_id}").status_code == 404
    actor["principal"] = Principal("hr", "tenant-b", frozenset({"HR"}))
    assert client.get(f"/api/cases/{case_id}").status_code == 404
    assert client.get("/api/hr/cases").json() == []
    app.dependency_overrides.pop(get_principal)
    assert client.get(f"/api/cases/{case_id}").status_code == 401


def test_http_review_validation_finalization(api):
    client, actor, engine = api
    case_id = create(client)
    document_ids = evidence(engine, case_id)
    payload = {
        "data": {
            "full_name": "Synthetic Person",
            "email": "synthetic@example.com",
            "phone": "+919876543210",
            "pan": "ABCDE1234F",
            "aadhaar": "234567890123",
        },
        "reviewed_document_ids": document_ids,
    }
    assert client.patch(f"/api/cases/{case_id}", json=payload).status_code == 403
    actor["principal"] = Principal("reviewer", "tenant-a", frozenset({"HR"}))
    patched = client.patch(f"/api/cases/{case_id}", json=payload)
    assert patched.status_code == 200, patched.text
    validated = client.post(f"/api/cases/{case_id}/validate")
    assert validated.status_code == 200, validated.text
    assert validated.json()["status"] == "validated", validated.text
    assert (
        client.post(
            f"/api/cases/{case_id}/finalize",
            json={"confirmed": False},
            headers={"Idempotency-Key": "test-key"},
        ).status_code
        == 422
    )
    first = client.post(
        f"/api/cases/{case_id}/finalize", json={"confirmed": True}, headers={"Idempotency-Key": "test-key"}
    )
    assert first.status_code == 200, first.text
    second = client.post(
        f"/api/cases/{case_id}/finalize", json={"confirmed": True}, headers={"Idempotency-Key": "test-key"}
    )
    assert second.status_code == 200
    assert first.json()["employee_id"] == second.json()["employee_id"]
    assert first.json()["status"] == "created"
    assert client.patch(f"/api/cases/{case_id}", json=payload).status_code == 409
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(Employee)) == 1


def test_http_hr_exports_filter_and_minimize_personal_data(api):
    client, actor, engine = api
    engineering_id = create(client)
    create(client, "Sales")
    assert client.get("/api/hr/export").status_code == 403
    with Session(engine) as db:
        case = db.get(Case, engineering_id)
        case.data = {
            "full_name": "=HYPERLINK(1)",
            "email": "synthetic@example.com",
            "pan": "ABCDE1234F",
            "aadhaar": "234567890123",
        }
        db.commit()
    actor["principal"] = Principal("outsider", "tenant-b", frozenset({"HR"}))
    create(client)
    actor["principal"] = Principal("reviewer", "tenant-a", frozenset({"HR"}))
    response = client.get("/api/hr/export", params={"department": "Engineering", "status": "received"})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 1
    assert rows[0]["case_id"] == engineering_id
    assert rows[0]["full_name"].startswith("'=")
    assert "pan" not in rows[0] and "aadhaar" not in rows[0]
    assert "ABCDE1234F" not in response.text
    xlsx = client.get("/api/hr/export", params={"format": "xlsx", "department": "Engineering"})
    assert xlsx.status_code == 200
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(xlsx.content), read_only=True).active
    assert sheet.max_row == 2
    assert sheet.cell(2, 5).value.startswith("'=")
