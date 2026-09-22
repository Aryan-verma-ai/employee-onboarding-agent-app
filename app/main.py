from contextlib import asynccontextmanager

from fastapi import Depends, FastAPI, Header
from pydantic import BaseModel, Field

from .auth import get_principal
from .config import settings, validate_configuration
from .confirmation import router as confirmation_router
from .dashboard import router as dashboard_router
from .db import engine, get_db
from .documents import router as documents_router
from .foundry import router as foundry_router
from .models import Base
from .service import OnboardingService
from .sso import router as sso_router


@asynccontextmanager
async def lifespan(app):
    validate_configuration()
    if settings.environment == "development":
        Base.metadata.create_all(engine)

    import os
    import threading
    import time
    from .auth import Principal
    from .worker import process_one

    def _worker_loop():
        tenant = os.getenv("WORKER_TENANT_ID", "a8780332-fecd-4f97-81b8-3782131af31a")
        subject = os.getenv("WORKER_SUBJECT", "onboarding-worker")
        principal = Principal(subject, tenant, frozenset({"HR"}))
        while True:
            try:
                processed = process_one(engine, principal)
                if not processed:
                    time.sleep(1.0)
            except Exception:
                time.sleep(2.0)

    worker_thread = threading.Thread(target=_worker_loop, daemon=True, name="extraction-worker-thread")
    worker_thread.start()
    yield


app = FastAPI(title="Foundry Employee Onboarding", lifespan=lifespan)


def service(db=Depends(get_db), principal=Depends(get_principal)):
    return OnboardingService(db, principal)


class CreateCase(BaseModel):
    department: str = Field(min_length=1, max_length=100)
    consent: bool


class Confirm(BaseModel):
    confirmed: bool


def case_response(case):
    full_name = (case.data or {}).get("full_name") or "Employee"
    start_date = (case.data or {}).get("start_date")
    onboarding_message = (
        f"🎉 Welcome aboard, {full_name}! Your onboarding is complete and employee profile is active. "
        f"Your official Employee ID is {case.employee_id}, and your assigned start date is {start_date}."
        if case.employee_id else None
    )
    return {
        "id": case.id,
        "department": case.department,
        "status": case.status,
        "data": case.data,
        "start_date": start_date,
        "missing_fields": case.missing_fields,
        "validation_outcomes": case.validation_outcomes,
        "consent_policy_version": case.consent_policy_version,
        "consent_withdrawn_at": case.consent_withdrawn_at,
        "employee_id": case.employee_id,
        "onboarding_message": onboarding_message,
        "created_at": case.created_at,
    }


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/cases", status_code=201)
def create_case(body: CreateCase, svc=Depends(service)):
    return case_response(svc.create_case(body.department, body.consent))


@app.get("/api/cases/{case_id}")
def get_case(case_id: str, svc=Depends(service)):
    return case_response(svc.get_case(case_id))


@app.post("/api/cases/{case_id}/validate")
def validate_case(case_id: str, svc=Depends(service)):
    case = svc.validate_case(case_id)
    if case.status == "created" or case.employee_id:
        from .notifications import send_welcome_email
        send_welcome_email(case.data, case.employee_id, case.department)
    return case_response(case)


@app.post("/api/cases/{case_id}/finalize")
def finalize_case(
    case_id: str, body: Confirm, idempotency_key: str = Header(alias="Idempotency-Key"), svc=Depends(service)
):
    case = svc.finalize_case(case_id, body.confirmed, idempotency_key)
    from .notifications import send_welcome_email
    send_welcome_email(case.data, case.employee_id, case.department)
    return case_response(case)


@app.post("/api/cases/{case_id}/email/welcome")
def send_welcome_email_route(case_id: str, svc=Depends(service)):
    case = svc.get_case(case_id)
    svc.require_consent(case)
    from .notifications import send_welcome_email
    res = send_welcome_email(case.data, case.employee_id or "PENDING", case.department)
    svc.audit(case.id, "email.welcome_sent", res)
    svc.db.commit()
    return res


app.include_router(documents_router)
app.include_router(dashboard_router)
app.include_router(foundry_router)
app.include_router(confirmation_router)
app.include_router(sso_router)


class UpdateCase(BaseModel):
    data: dict[str, str] = Field(default_factory=dict)
    reviewed_document_ids: list[str] = Field(default_factory=list)


@app.patch("/api/cases/{case_id}")
def update_case(case_id: str, body: UpdateCase, svc=Depends(service)):
    from fastapi import HTTPException
    from sqlalchemy import select

    from .models import Document
    from .validation import ALLOWED_FIELDS, REQUIRED_FIELDS

    case = svc.get_case(case_id)
    svc.require_consent(case)
    case.validation_outcomes = []
    if case.status == "created":
        raise HTTPException(409, "Completed cases are immutable")
    if set(body.data) - set(ALLOWED_FIELDS) or any(len(value) > 500 for value in body.data.values()):
        raise HTTPException(422, "Unsupported field or excessive field length")
    if body.reviewed_document_ids and not svc.principal.is_hr:
        raise HTTPException(403, "Only HR can attest document review")
    docs = []
    for document_id in set(body.reviewed_document_ids):
        doc = svc.db.scalar(
            select(Document).where(
                Document.id == document_id, Document.case_id == case.id, Document.tenant_id == case.tenant_id
            )
        )
        if not doc or doc.extraction.get("status") != "complete":
            raise HTTPException(
                422, "Reviewed document must belong to this case and have completed extraction"
            )
        docs.append(doc)
    if any(case.data.get(key) != value for key, value in body.data.items()):
        # Corrections invalidate prior attestations; HR must review new values.
        for old in svc.db.scalars(select(Document).where(Document.case_id == case.id)):
            old.extraction = {**old.extraction, "reviewed": False}
    case.data = {**case.data, **body.data}
    for doc in docs:
        doc.extraction = {**doc.extraction, "reviewed": True}
    if case.status == "failed":
        svc.transition(case, "extracting")
    svc.transition(case, "needs-information")
    svc.audit(
        case.id,
        "record-corrected",
        {"fields": sorted(body.data), "reviewed_document_ids": sorted(body.reviewed_document_ids)},
    )
    svc.db.commit()
    return case_response(case)


@app.get("/ready")
def ready():
    from fastapi import HTTPException
    from sqlalchemy import text
    from sqlalchemy.exc import SQLAlchemyError

    try:
        validate_configuration()
        with engine.connect() as connection:
            connection.execute(
                text(
                    "SELECT consent_withdrawn_at, consent_policy_version, validation_outcomes FROM onboarding_cases LIMIT 0"
                )
            )
            connection.execute(text("SELECT lease_token, lease_until FROM extraction_jobs LIMIT 0"))
    except (RuntimeError, SQLAlchemyError):
        raise HTTPException(503, "Service configuration or database migration is not ready") from None
    return {"status": "ready", "cloud_connectivity": "not_probed"}


@app.post("/api/cases/{case_id}/consent/withdraw")
def withdraw_consent(case_id: str, body: Confirm, svc=Depends(service)):
    from datetime import datetime, timezone

    from fastapi import HTTPException

    case = svc.get_case(case_id)
    if svc.principal.subject != case.owner_id:
        raise HTTPException(403, "Only the employee can withdraw their consent")
    if body.confirmed is not True:
        raise HTTPException(422, "Explicit withdrawal confirmation required")
    if not case.consent_withdrawn_at:
        case.consent_withdrawn_at = datetime.now(timezone.utc)
        case.validation_outcomes = []
        svc.audit(case.id, "consent-withdrawn", {"policy_version": case.consent_policy_version})
        svc.db.commit()
    return case_response(case)


@app.get("/api/cases/{case_id}/profile/export")
def export_employee_profile(case_id: str, svc=Depends(service)):
    """Generate an Excel employee profile from the case data.

    Available at any stage — exports whatever data has been collected so far.
    After employee creation the employee ID is included.
    """
    import io

    from fastapi import HTTPException
    from fastapi.responses import Response

    case = svc.get_case(case_id)
    svc.require_consent(case)

    data = case.data or {}
    if not data:
        raise HTTPException(422, "No employee data to export yet")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Profile"

    # ── Styles ──
    header_font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid")
    label_font = Font(name="Calibri", bold=True, size=11)
    value_font = Font(name="Calibri", size=11)
    label_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ── Title Row ──
    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = "Employee Onboarding Profile"
    title_cell.font = header_font
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws["B1"].fill = header_fill
    ws.row_dimensions[1].height = 35

    # ── Column widths ──
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 45

    # ── Profile fields ──
    profile_fields = [
        ("Employee ID", case.employee_id or "Pending creation"),
        ("Case ID", case.id),
        ("Department", case.department.upper() if case.department else ""),
        ("Onboarding Status", "Active / Created" if case.status == "created" else case.status.replace("-", " ").title()),
        ("Start Working Date", data.get("start_date", "Pending finalization")),
        ("", ""),  # spacer
        ("Full Name", data.get("full_name", "")),
        ("Email", data.get("email", "")),
        ("Phone", data.get("phone", "")),
        ("PAN Number", data.get("pan", "")),
        ("Aadhaar Number", data.get("aadhaar", "")),
        ("Date of Birth", data.get("dob", "")),
        ("Address", data.get("address", "")),
        ("", ""),  # spacer
        ("Created At", case.created_at.strftime("%Y-%m-%d %H:%M UTC") if case.created_at else ""),
    ]

    for i, (label, value) in enumerate(profile_fields, start=2):
        label_cell = ws.cell(row=i, column=1, value=label)
        value_cell = ws.cell(row=i, column=2, value=value)
        if label:
            label_cell.font = label_font
            label_cell.fill = label_fill
            label_cell.border = thin_border
            value_cell.font = value_font
            value_cell.border = thin_border
            value_cell.alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    svc.audit(case.id, "profile.exported", {"format": "xlsx"})
    svc.db.commit()

    filename = f"employee_profile_{data.get('full_name', 'unknown').replace(' ', '_')}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
